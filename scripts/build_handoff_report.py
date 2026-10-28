"""Rebuild the Word handoff using actual delivery manifests and validation records."""
from __future__ import annotations
import argparse
import json
from pathlib import Path
import sys
from datetime import datetime

ROOT=Path(__file__).resolve().parents[1]


def read(path,default=None):
    target=ROOT/path
    return json.loads(target.read_text(encoding='utf-8')) if target.exists() else default


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'docs/ArtifactHealthRecord_xqw-v0_交付与后续工作.docx')
    args=parser.parse_args()
    from docx import Document
    from docx.image.image import Image as DocxImage
    from docx.image.exceptions import UnrecognizedImageError
    from docx.shared import Inches,Pt,RGBColor,Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE

    manifest=read(Path('data/public/mural_reference/manifest.json'),{})
    validation=read(Path('docs/xqw-v0-validation.json'),{})
    models=read(Path('ai/REAL_ASSET_VALIDATION.json'),{})
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    for name in ['Normal','Body Text','Table Grid','List Bullet','List Number']:
        style=doc.styles[name]
        style.font.name='Calibri';style.font.size=Pt(10)
        style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        style.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',27),('Heading 1',17),('Heading 2',12)]:
        style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string('274C4C')
        style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('ArtifactHealthRecord · xqw-v0  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    section.header.paragraphs[0].text='文物病历  /  项目书对照与后续实施清单'

    def p(text='',bold=False):
        para=doc.add_paragraph();run=para.add_run(text);run.bold=bold;return para
    def heading(text,level=1):doc.add_heading(text,level)
    def page():doc.add_page_break()
    def bullet(text):doc.add_paragraph(text,style='List Bullet')
    def link(label,url):
        para=doc.add_paragraph()
        relation=para.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
        el=OxmlElement('w:hyperlink');el.set(qn('r:id'),relation)
        run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'167175');pr.append(color);run.append(pr)
        text=OxmlElement('w:t');text.text=label+'  '+url;run.append(text);el.append(run);para._p.append(el)
    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.autofit=False
        widths={2:[8.7,8.7],3:[3.2,5.4,8.8],4:[2.0,3.0,3.0,9.4]}.get(len(headers),[17.4/len(headers)]*len(headers))
        for column,width in zip(t.columns,widths):column.width=Cm(width)
        for cell,text,width in zip(t.rows[0].cells,headers,widths):
            cell.width=Cm(width)
            cell.text=text
            for run in cell.paragraphs[0].runs:run.bold=True
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8F0EE');cell._tc.get_or_add_tcPr().append(shade)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for cell,text,width in zip(cells,row,widths):
                cell.width=Cm(width)
                cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(3)
                    for run in para.runs:run.font.size=Pt(9)
            prevent=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(prevent)
        p()
        return t

    doc.add_heading('文物病历',0)
    p('xqw-v0 交付与后续工作',True)
    p('对照《方向一_文物病历_比赛构建指南_V1.0》')
    p('核查日期：2026年10月28日（Asia/Shanghai）')
    p('工程范围：ArtifactHealthRecord；保留完整业务与AI工程。')
    p('目前最需要补的，是同一目标地域的真实病害证据和专业验证。工程可以登记、辅助标注、编方案、会诊和归档，但公开馆藏照片不能代替专家标注，完整代码不能代替真实试点。',True)
    table(['这次交付','仍缺的关键实物/证据'],[
        ['完整前后端、病历—病例—方案—规则—会诊—档案闭环','目标地域30–50张局部图及配套采集记录'],
        [f"{manifest.get('count',0)}张馆方开放照片、{manifest.get('artifact_count',0)}件对象、原件和许可元数据",'六类专家标注；完整授权修复病例20份'],
        ['来源复核、SAM交互、标准图符执行、数据与发布验证','真实训练/独立评测、图符专业验收'],
        ['本地工程测试与可复现证据，GitHub指定分支','华为云/CodeArts真实记录、机构试点、比赛材料'],
    ])
    p('阅读顺序：先看第1–3部分了解图片到底是什么，再按第4部分领取任务；第5–7部分给出验收口径、工程验证与复现入口。')
    link('代码分支','https://github.com/Xlovlhj/ArtifactHealthRecord/tree/xqw-v0')
    p('数量来自公开数据manifest，测试来自本轮validation；未取得的指标不写成已完成。')

    page();heading('1. 现在的图片是什么，能拿它做什么')
    counts=manifest.get('view_counts',{})
    p(f"公开包实际有{manifest.get('count',0)}张照片，对应{manifest.get('artifact_count',0)}件馆藏对象。其中彩色正面{counts.get('front',0)}张，黑白正面{counts.get('front_monochrome',0)}张，背面/装裱结构{counts.get('reverse',0)}张。同一物件的不同视图不是新的独立壁画。一张线描说明图已人工排除。")
    p('这些是博物馆公开的真实中国壁画照片，包含山西、新疆等。它们能够让工作台使用真实文物图像完成查看、来源登记、标注和检索演练；不能据此说已收集了甘肃单地域现场病害数据。')
    table(['资产','数量/状态','现在能用在哪'],[
        ['官方开放馆藏照片',f"{manifest.get('count',0)}张 / {manifest.get('artifact_count',0)}件；随仓库提供",'真实图像展示、标注操作、通用模型工程检查；不作六类监督真值'],
        ['已有敦煌研究GT裁片','本机60张；公开仓库不分发','镜像许可声明不足以替代完整权利链；同源裁片保守同组，不能随机分训练测试'],
        ['MuralDH完整数据','0份；保存了元数据及下载失败记录','正常取得授权下载后再核尺寸/哈希、类别映射和原址分组'],
        ['程序绘制教学样板','3份，继续保留且标合成','软件回归、流程演示、合成训练管线检查；不是实拍'],
        ['六类病害专家真值','0张','必须真实采集、双人标注、第三人仲裁后建立'],
        ['完整授权修复病例','0份；现有4条甘肃公开摘要','摘要可查资料，不能凑成项目书要求的20份完整病例'],
    ])
    p('公开包逐图保存对象编号、馆方链接、JPEG原件、尺寸、SHA-256、视图类型、材料/年代原字段、许可证据和分组依据。18份官方对象API元数据可回查isPublicDomain字段。代码MIT许可证不扩大外部数据授权。')
    link('公开数据清单','https://github.com/Xlovlhj/ArtifactHealthRecord/blob/xqw-v0/data/public/mural_reference/manifest.json')
    link('The Met开放政策','https://www.metmuseum.org/hubs/open-access')
    link('MuralDH数据发布页','https://datadryad.org/dataset/doi:10.5061/dryad.bnzs7h4jd')

    page();heading('2. 真实照片样例与可追溯出处')
    p('以下直接使用交付包中的馆方照片；没有绘制病害标签，也没有生成“修复后”画面。图片只是资料样例，是否存在某类病害仍待专业判断。')
    front=[r for r in manifest.get('records',[]) if r.get('view_type')=='front']
    chosen=[]
    for row in front:
        # Some valid museum JPEGs use header variants unsupported by python-docx.
        # Select supported originals rather than altering the evidence files.
        try:
            DocxImage.from_file(str(ROOT/'data/public/mural_reference'/row['path']))
        except UnrecognizedImageError:
            continue
        if row['artifact_id'] not in {r['artifact_id'] for r in chosen}:
            chosen.append(row)
        if len(chosen)==2:break
    for row in chosen:
        heading(row.get('title_zh') or row['title'],2)
        path=ROOT/'data/public/mural_reference'/row['path']
        # Scale only for the page; original source bytes remain in the public package.
        max_w,max_h=6.0,2.3
        width=min(max_w,max_h*row['width']/row['height'])
        doc.add_picture(str(path),width=Inches(width))
        p(f"对象：{row['artifact_id']}；图像：{row['id']}；地区：{row['region']}；尺寸：{row['width']}×{row['height']}；许可：{row['license']}。")
        p('来源：'+row['source_url'])
        p('原件SHA-256：'+row['sha256'])

    page();heading('3. 本次改了什么，原来的工程保留了什么')
    p('原项目已经有大量业务和AI代码，但存在来源登记不足、说明文档过期、图符数据没有接到渲染、SAM缺业务入口等问题。这次按实际调用链补齐，没有删成简单档案管理页。')
    table(['范围','本次工作','原有完整能力'],[
        ['真实数据','收集官方开放照片、核图、保存原件/元数据、独立参考图库、显式导入','保留原有研究数据清单、下载器与教学测试'],
        ['影像来源','逐图来源/许可/用途/分组、录入人与复核人分离、来源变化撤销复核、真实/合成计数','保留原始字节、方向规范化显示图、标尺与面积'],
        ['图符和标注','接入六类图符参数；裂隙中心线独立录入；SAM点选、候选预览和人工采用','保留多边形编辑、六类标注、SVG/PDF、模型候选'],
        ['病例与证据','编辑保留附件，补独立复核；公开摘要/参考图与完整病例分开','保留BM25、可选embedding、DINOv2/CLIP检索与证据快照'],
        ['训练与数据质控','NPZ原始多类标注纳入哈希校验；采集模板、公开包完整性检查','保留U-Net/ResNet34/SegFormer、二分类/六类、迁移、冻结、五折与失败样本'],
        ['环境与交付','安装完整固定AI依赖；离线verify；公开包过滤；Word与实际验收记录','保留RDS/Redis/OBS/MaaS适配、容器配置、Python启动入口'],
        ['核心工作流','逐项回归与界面引导','13节方案、材料试验、预算、规则、独立会诊、敏感导出、10类档案、审计与迁移'],
    ])
    p('没有更换原有病害算法或改写损失权重。SAM仅沿用既有本地通用模型提边界；图符遵循已登记的参数包；新增发布与数据检查不替代专家决策。')
    p('图符资料核验、程序绘制、专家成图验收是三件事。当前保留rendering_certified=false；缺失裂隙中心线时明确提示，不随意猜出一条中心线。')

    page();heading('4. 按先后顺序还要完成的工作')
    p('下表的“完成”都要求能拿出文件、原始记录或可复现结果。责任写角色，具体姓名由团队分工确定；没有假填合作单位或专家。')
    tasks=[
        ['P0 / 先做','目标地域与合作','负责人+文保单位','落实1家能供真实资料的单位、1位能审核的专家；最终扩展3家试点。取得用途/公开范围/保密说明，确定甘肃具体场景。'],
        ['P0 / 同步','现场影像与采集','采集负责人','30–50张真实局部图起步；保留原始文件、对象/位置对应、时间、设备、照明、标尺/色卡、多角度和质量记录。至少覆盖六类，不用跨地域馆藏图凑数。'],
        ['P0 / 影像后','六类标注','2名标注者+仲裁专家','起甲、酥碱、空鼓、裂隙、颜料层脱落、地仗脱落。初赛每类≥20实例，六类合计≥120；后续按指南质控目标每类100。保留人员、日期、版本、争议和一致性。'],
        ['P0 / 专项','现场病害证据','文保专家+检测人员','空鼓需要现场调查/检测；酥碱需要相应盐分/材料等证据。将检测报告与区域关联，未知不标成已确诊。'],
        ['P0 / 数据后','独立划分与模型','CV负责人','按原壁画/遗址先分组再切片，隔离训练验证测试；同组多视图不能跨集合。对比U-Net/SegFormer，记录数据/权重SHA、每类指标和失败案例。'],
        ['P0 / 并行','完整病例20份','资料负责人+专家','每例含对象、病害依据、环境/材料、处理、材料试验、结果/随访、失败与限制、来源和授权。录入并由另一账号复核；新闻摘要不是完整病例。'],
        ['P0 / 图件','标准成图验收','制图人员+文保专家','对六类符号、纸面线宽/间距、图例、裂隙中心线、尺度与面积逐项检查；保存标准版本、检查表、意见和返工记录。'],
        ['P0 / 比赛链路','CodeArts与华为云','后端/云负责人','真实使用CodeArts并留任务/构建/修复证据；接入RDS、Redis、OBS、MaaS，验证上传、缓存、生成、权限、备份和恢复；部署HTTPS Demo。'],
        ['P1 / 闭环','模拟与前后对比','CV+文保专家','取得适用LaMa/SD权重授权并实测，保留水印/不确定区域；收集同条件术前术后照片，不能把像素变化当修复成功。'],
        ['P1 / 实地','3家试点与用户研究','负责人+试点联系人','以3家、10位修复师、30件文物为项目书目标，采集真实人工/系统计时、返工和SUS原始问卷，保留参与/数据授权。'],
        ['P1 / 有数据后','E1–E6完整评估','评测负责人+独立专家','按第5部分执行对照/盲评、真实采样、置信区间和失败分析。没有数据不填成绩。'],
        ['P1 / 提交前','比赛材料','队长+材料负责人','当届官方模板1、PPT、3分钟视频、公开Demo与开源截图、人员分工、合作证明和三次彩排。按校内/官方通知复核截止日期。'],
    ]
    table(['顺序','事项','责任角色','明确完成标准'],tasks)
    p('建议执行次序：先确定数据与专家，再完成采集和病例；标注与云链路可并行；真实训练和盲评依赖数据完成。仅继续美化界面无法补掉当前主要差距。')

    page();heading('5. 怎么证明做完：实验与验收证据')
    table(['实验','按项目书准备什么','输出与不可混淆的边界'],[
        ['E1 分割','30–50件/100–200图目标、五折按原壁画或遗址分组；U-Net/SegFormer等基线','每类Dice/IoU/Precision/Recall、宏平均±标准差、混淆矩阵、最差20样本；SAM通用推理成功不算六类Dice。'],
        ['E2 病害图','30份图、3位专家；人工/模板/系统对照','按固定检查表报合规率、返工数；图符JSON的verified不等于专家合规率。'],
        ['E3 字段抽取','30份真实病历；人工/VLM/VLM+Schema','逐字段金标准、准确率、完整率；未知值不当正确负例，结构合法不等于事实准确。'],
        ['E4 病例检索','50条查询；CLIP/DINOv2/过滤对照，专家相关性标注','Hit@5、Precision@5、MRR和相关性；视觉相似度不是临床相似概率。'],
        ['E5 方案质量','10病例×3方法；3专家盲评LLM/RAG/RAG+规则','完整性、依据、风险、可执行性1–5分；引文逐句核支持关系、采纳率、失败原因。'],
        ['E6 端到端','3家、10位修复师、30件；人工/系统配对','单件实际计时、返工、SUS十题原始答卷；不能把项目书“4小时→30分钟”当实测。'],
    ])
    p('每个实验保留采样规则、原始数据、人员/盲法、版本、运行命令、分母、排除比例、全部结果与失败样本。没有相关类别或没有真值时指标为null，不记满分。')
    p('初赛门槛例子Dice≥0.65、后续宏平均≥0.75、SUS≥68均来自项目书的验收目标，当前没有这些成绩。是否可用于具体文物，还须专家结合漏检、误检和适用范围判断。')
    heading('可直接用的采集与评测入口',2)
    bullet('data/templates/：真实采集、标注审核等登记模板；字段留空表示尚未取得，不会自动生成专家意见。')
    bullet('scripts/prepare_dataset.py：LabelMe→掩膜，来源分组、原图/标注哈希；scripts/train.py、scripts/evaluate.py：正式训练和E1。')
    bullet('scripts/evaluate_study.py --init E2 --output runtime/studies/E2.json：生成空采集表；E3至E6同理。录入真实观察后用--input统计。')
    bullet('标准图件、来源复核、会诊意见与正式验收需要各自留痕；同一账号不能替代独立专家复核。')

    page();heading('6. 本轮验证到了哪一步')
    p('工程验证回答“这条代码与操作能否正确执行”，专业实验回答“它对真实文物是否有用”。以下只报告已经留有记录的工程与通用模型验证。')
    if validation:
        summary=validation.get('summary',{})
        table(['验证项','实际结果'],[
            ['Python回归',summary.get('pytest','见xqw-v0-validation.json')],
            ['前端构建',summary.get('frontend_build','见本轮记录')],
            ['浏览器完整工作流',summary.get('browser','见本轮记录')],
            ['公开数据与干净副本',summary.get('delivery','见本轮记录')],
            ['实际模型',summary.get('models','见../ai/REAL_ASSET_VALIDATION.json')],
        ])
    else:
        p('本轮验证记录尚未生成；此Word为工作稿，完成验证后须重新生成。',True)
    p('验证范围包括来源复核和撤销、跨机构隔离、版本冲突、SAM候选绑定原图哈希、无权重返回、图符导出、真实图库导入、方案/会诊、敏感导出和原始档案。浏览器里的候选UI夹具只用于测试交互，真实模型执行另有证据。')
    p('当前真实六类Dice、真实病例检索相关性、方案采纳率、SUS均无结果；云服务适配器的离线测试不等于实际华为云连通。Docker/新机器安装、PostgreSQL并发与正式备份恢复仍需目标环境验收。')
    heading('工程执行记录',2)
    for name in ['docs/xqw-v0-validation.json','docs/frontend-validation.json','ai/REAL_ASSET_VALIDATION.json','docs/guide-implementation-matrix.md']:
        p(name)
    p('早期validation文件为历史记录，能力与数量以本次xqw-v0记录为准。不能沿用早期README的旧数字或“功能尚无”等过时说明。')

    page();heading('7. 接手运行与下一次交付')
    p('本机启动：python3 /home/xqw/xqw-file/HRF/ArtifactHealthRecord/run.py')
    p('新机器：安装Miniconda，进入项目目录，执行：')
    for command in ['python3 run.py setup --with-ai','python3 run.py verify','python3 run.py']:
        p(command,True)
    p('浏览器打开 http://127.0.0.1:8765。演示账号restorer/expert等，密码ahr-demo-2026；正式使用独立数据库、强密码和生产配置。所有启动逻辑用Python，没有.sh启动脚本。')
    p('公开参考图库已包含图片原件，不依赖打开页面时外网下载。通用权重不随Git上传，可用python3 run.py script fetch_models.py --models dinov2 sam resnet34显式下载并校验。真实六类权重必须另行训练。')
    p('下一次交付至少要增加：真实采集包与授权、六类专家标注/仲裁、20份完整病例、首次独立分割评估、图件专家核查、实际云链路与CodeArts证据。届时重新生成数据清单和本报告，不能只改README里的完成状态。')
    table(['应保存的文件','为什么需要'],[
        ['原图+逐图来源/授权/分组+色卡/标尺/检测','能确认图片身份、用途与测量依据'],
        ['双人LabelMe、仲裁记录、掩膜清单','能解释病害标签如何得出、有没有跨组泄漏'],
        ['病例证据、处理与随访、专家复核','能支持方案推理，不靠文字摘要凑数'],
        ['数据/模型哈希、命令、每折结果和失败样本','能由第三方复现真实效果'],
        ['云部署/CodeArts/备份恢复/访问日志','证明比赛技术底座真实使用'],
        ['计时/SUS/盲评原始表、单位与专家证明','证明有人试用以及收益是否成立'],
    ])
    heading('资料来源与使用说明',2)
    p('主要需求依据：用户提供的Plan/方向一_文物病历_比赛构建指南_V1.0.pdf；本报告只在项目内保存逐项对照，不重新分发项目书全文。')
    p('标准依据与来源页登记见data/knowledge/manifest.json、symbol_package_gbt30237_2025.json及docs/standards-verification.md。报告不将项目内部复核解释为行政审批或正式资质。')
    p('照片来自The Met官方对象记录与开放政策；遵从CC0授权并保留出处。第三方模型、代码与研究数据分别遵循上游条件，见THIRD_PARTY_NOTICES.md。')
    link('公开交付分支','https://github.com/Xlovlhj/ArtifactHealthRecord/tree/xqw-v0')
    args.output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(args.output)
    print(json.dumps({'output':str(args.output),'image_count':manifest.get('count'),'artifact_count':manifest.get('artifact_count'),
                      'view_counts':manifest.get('view_counts'),'validation_record_present':bool(validation)},ensure_ascii=False))


if __name__=='__main__':main()
